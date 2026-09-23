#!/usr/bin/env python3
"""Bounded real QPS RTM workload: source-backed control snapshot -> Excel -> semantic receipt.

This intentionally reproduces the expected v0.6 calculation from repository
control data. It does not regenerate the authoritative v0.5 workbook and must
not promote the expected counts into contract, compliance, engineering, or
negotiation authority.
"""
from __future__ import annotations

import csv
import hashlib
import json
from collections import Counter
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


BASE = Path("data/qps_rtm_partial_relax_v06")
BASELINE = BASE / "baseline_v05.json"
BINDINGS = BASE / "bindings.csv"
EXPECTED = BASE / "expected_v06.json"
OUT = Path("artifacts/qps_rtm_workload")


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def loads_json_lossless(text: str) -> dict[str, Any]:
    return json.loads(text, parse_float=Decimal)


def load_inputs() -> tuple[dict[str, Any], list[dict[str, str]], dict[str, Any]]:
    baseline = loads_json_lossless(BASELINE.read_text(encoding="utf-8"))
    expected = loads_json_lossless(EXPECTED.read_text(encoding="utf-8"))
    with BINDINGS.open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return baseline, rows, expected


def is_numeric_zero(value: Any) -> bool:
    if type(value) is int:
        return value == 0
    if isinstance(value, Decimal):
        return value.is_finite() and value == Decimal(0)
    return False


def validate_authority_guards(
    baseline: dict[str, Any],
    expected: dict[str, Any],
) -> dict[str, Any]:
    required_promotion_gate = (
        "exact source v0.5 binary must be regenerated and the emitted receipt must PASS "
        "before these become the current workbook counts"
    )
    failures: list[str] = []

    if baseline.get("authority_transfer") is not False:
        failures.append("baseline authority_transfer must be false")
    if not is_numeric_zero(baseline.get("formal_credit_delta")):
        failures.append("baseline formal_credit_delta must be numeric zero, not bool")
    if baseline.get("status") != "PASS":
        failures.append("baseline validation receipt status must be PASS")

    if expected.get("state") != "EXPECTED_NOT_YET_CREDITED":
        failures.append("expected state must remain EXPECTED_NOT_YET_CREDITED")
    if expected.get("authority_transfer") is not False:
        failures.append("expected authority_transfer must be false")
    for key in (
        "formal_credit_delta",
        "engineering_credit_delta",
        "negotiation_credit_delta",
        "compliance_credit_delta",
    ):
        if not is_numeric_zero(expected.get(key)):
            failures.append(f"{key} must be numeric zero, not bool")
    if expected.get("promotion_gate") != required_promotion_gate:
        failures.append("exact-v0.5 regeneration promotion gate changed or weakened")

    if failures:
        raise ValueError("non-compensating authority guard failure: " + "; ".join(failures))

    return {
        "all_non_compensating_guards_passed": True,
        "baseline_authority_transfer": baseline["authority_transfer"],
        "baseline_formal_credit_delta": baseline["formal_credit_delta"],
        "expected_state": expected["state"],
        "expected_authority_transfer": expected["authority_transfer"],
        "formal_credit_delta": expected["formal_credit_delta"],
        "engineering_credit_delta": expected["engineering_credit_delta"],
        "negotiation_credit_delta": expected["negotiation_credit_delta"],
        "compliance_credit_delta": expected["compliance_credit_delta"],
        "promotion_gate": expected["promotion_gate"],
    }


def validate_bindings(rows: list[dict[str, str]]) -> None:
    if len(rows) != 5:
        raise ValueError(f"expected 5 exact bindings, got {len(rows)}")
    if len({r["native_id"] for r in rows}) != 5:
        raise ValueError("native_id values are not unique")
    if len({r["canonical_target"] for r in rows}) != 5:
        raise ValueError("canonical_target values are not unique")
    if any(r["exact_binding"].strip().lower() != "true" for r in rows):
        raise ValueError("all workload rows must be exact_binding=true")

    required = {
        "TEC_ID_034": ("RTM-138", "LKT_ONLY"),
        "TEC_ID_205": ("RTM-691", "LKT_ONLY"),
        "TEC_ID_206": ("RTM-693", "BOTH_CHALLENGE_SAME_ATOM"),
        "TEC_ID_207": ("RTM-695", "BOTH_CHALLENGE_SAME_ATOM"),
        "TEC_ID_144": ("RTM-521", "BOTH_CHALLENGE_SAME_ATOM"),
    }
    for row in rows:
        target, queue = required[row["native_id"]]
        if row["canonical_target"] != target or row["queue_after_regeneration"] != queue:
            raise ValueError(f"binding mismatch for {row['native_id']}")


def calculate(
    baseline: dict[str, Any],
    rows: list[dict[str, str]],
) -> dict[str, Any]:
    peer = dict(baseline["peer_queues"])
    increments = Counter(r["queue_after_regeneration"] for r in rows)
    for queue, delta in increments.items():
        peer[queue] = peer.get(queue, 0) + delta

    coverage = dict(baseline["coverage_states"])
    coverage["ATOMIZED_SOURCE_BACKED"] += len(rows)
    coverage["LKT_SECTION_FAMILY_ALIGNMENT_PENDING"] -= len(rows)

    return {
        "canonical_rtm_denominator": baseline["rtm_denominator"],
        "atomic_rows": baseline["atomic_rows"] + len(rows),
        "peer_queues": peer,
        "parent_screening": coverage,
        "ranked_v05_family_frontier_remaining": 0,
    }


def assert_expected(calculated: dict[str, Any], expected: dict[str, Any]) -> None:
    compare = {
        "canonical_rtm_denominator": expected["canonical_rtm_denominator"],
        "atomic_rows": expected["atomic_rows"],
        "peer_queues": expected["peer_queues"],
        "parent_screening": expected["parent_screening"],
        "ranked_v05_family_frontier_remaining": expected["ranked_v05_family_frontier_remaining"],
    }
    if calculated != compare:
        raise AssertionError(
            "calculated workload output differs from repository expected control: "
            + json.dumps({"calculated": calculated, "expected": compare}, sort_keys=True)
        )


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_workbook(
    baseline: dict[str, Any],
    rows: list[dict[str, str]],
    expected: dict[str, Any],
    calculated: dict[str, Any],
    path: Path,
) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "BINDINGS"

    binding_headers = list(rows[0].keys())
    ws.append(binding_headers)
    for row in rows:
        ws.append([row[h] for h in binding_headers])
    style_header(ws)
    for col in "ABCDEFGHI":
        ws.column_dimensions[col].width = 22
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["D"].width = 80
    ws.column_dimensions["E"].width = 34
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    base_ws = wb.create_sheet("BASELINE_V05")
    base_ws.append(["metric", "key", "value"])
    base_rows = [
        ("scalar", "rtm_denominator", baseline["rtm_denominator"]),
        ("scalar", "atomic_rows", baseline["atomic_rows"]),
        ("peer_queue", "ALAT_ONLY", baseline["peer_queues"]["ALAT_ONLY"]),
        ("peer_queue", "BOTH_CHALLENGE_SAME_ATOM", baseline["peer_queues"]["BOTH_CHALLENGE_SAME_ATOM"]),
        ("peer_queue", "LKT_ONLY", baseline["peer_queues"]["LKT_ONLY"]),
        ("peer_queue", "PROTECTED_REMAINDER", baseline["peer_queues"]["PROTECTED_REMAINDER"]),
        ("parent_screening", "ATOMIZED_SOURCE_BACKED", baseline["coverage_states"]["ATOMIZED_SOURCE_BACKED"]),
        ("parent_screening", "LKT_SCOPE_INTERPRETATION_BOUND", baseline["coverage_states"]["LKT_SCOPE_INTERPRETATION_BOUND"]),
        ("parent_screening", "LKT_SECTION_FAMILY_ALIGNMENT_PENDING", baseline["coverage_states"]["LKT_SECTION_FAMILY_ALIGNMENT_PENDING"]),
        ("parent_screening", "SOURCE_EXTRACTION_REQUIRED", baseline["coverage_states"]["SOURCE_EXTRACTION_REQUIRED"]),
    ]
    for row in base_rows:
        base_ws.append(list(row))
    style_header(base_ws)

    calc_ws = wb.create_sheet("CALCULATED_V06")
    calc_ws.append(["metric", "key", "formula", "expected"])
    formula_rows = [
        ("scalar", "canonical_rtm_denominator", "=BASELINE_V05!C2", expected["canonical_rtm_denominator"]),
        ("scalar", "atomic_rows", "=BASELINE_V05!C3+COUNTA(BINDINGS!A2:A6)", expected["atomic_rows"]),
        ("peer_queue", "ALAT_ONLY", '=BASELINE_V05!C4+COUNTIF(BINDINGS!H2:H6,"ALAT_ONLY")', expected["peer_queues"]["ALAT_ONLY"]),
        ("peer_queue", "BOTH_CHALLENGE_SAME_ATOM", '=BASELINE_V05!C5+COUNTIF(BINDINGS!H2:H6,"BOTH_CHALLENGE_SAME_ATOM")', expected["peer_queues"]["BOTH_CHALLENGE_SAME_ATOM"]),
        ("peer_queue", "LKT_ONLY", '=BASELINE_V05!C6+COUNTIF(BINDINGS!H2:H6,"LKT_ONLY")', expected["peer_queues"]["LKT_ONLY"]),
        ("peer_queue", "PROTECTED_REMAINDER", "=BASELINE_V05!C7", expected["peer_queues"]["PROTECTED_REMAINDER"]),
        ("parent_screening", "ATOMIZED_SOURCE_BACKED", "=BASELINE_V05!C8+COUNTA(BINDINGS!A2:A6)", expected["parent_screening"]["ATOMIZED_SOURCE_BACKED"]),
        ("parent_screening", "LKT_SCOPE_INTERPRETATION_BOUND", "=BASELINE_V05!C9", expected["parent_screening"]["LKT_SCOPE_INTERPRETATION_BOUND"]),
        ("parent_screening", "LKT_SECTION_FAMILY_ALIGNMENT_PENDING", "=BASELINE_V05!C10-COUNTA(BINDINGS!A2:A6)", expected["parent_screening"]["LKT_SECTION_FAMILY_ALIGNMENT_PENDING"]),
        ("parent_screening", "SOURCE_EXTRACTION_REQUIRED", "=BASELINE_V05!C11", expected["parent_screening"]["SOURCE_EXTRACTION_REQUIRED"]),
        ("scalar", "ranked_v05_family_frontier_remaining", "=0", expected["ranked_v05_family_frontier_remaining"]),
    ]
    for row in formula_rows:
        calc_ws.append(list(row))
    style_header(calc_ws)
    calc_ws.column_dimensions["B"].width = 42
    calc_ws.column_dimensions["C"].width = 58

    check_ws = wb.create_sheet("PYTHON_CHECK")
    check_ws.append(["metric", "key", "calculated_value"])
    check_ws.append(["scalar", "canonical_rtm_denominator", calculated["canonical_rtm_denominator"]])
    check_ws.append(["scalar", "atomic_rows", calculated["atomic_rows"]])
    for key in ("ALAT_ONLY", "BOTH_CHALLENGE_SAME_ATOM", "LKT_ONLY", "PROTECTED_REMAINDER"):
        check_ws.append(["peer_queue", key, calculated["peer_queues"][key]])
    for key in (
        "ATOMIZED_SOURCE_BACKED",
        "LKT_SCOPE_INTERPRETATION_BOUND",
        "LKT_SECTION_FAMILY_ALIGNMENT_PENDING",
        "SOURCE_EXTRACTION_REQUIRED",
    ):
        check_ws.append(["parent_screening", key, calculated["parent_screening"][key]])
    check_ws.append(["scalar", "ranked_v05_family_frontier_remaining", calculated["ranked_v05_family_frontier_remaining"]])
    style_header(check_ws)

    prov_ws = wb.create_sheet("PROVENANCE")
    prov_ws.append(["field", "value"])
    provenance = [
        ("baseline_source_repository", baseline["source_repository"]),
        ("baseline_source_commit", baseline["source_commit"]),
        ("baseline_source_path", baseline["source_path"]),
        ("baseline_source_blob_sha", baseline["source_blob_sha"]),
        ("baseline_source_url", f"https://github.com/{baseline['source_repository']}/blob/{baseline['source_commit']}/{baseline['source_path']}"),
        ("expected_source_repository", expected["source_repository"]),
        ("expected_source_commit", expected["source_commit"]),
        ("expected_source_path", expected["source_path"]),
        ("expected_source_blob_sha", expected["source_blob_sha"]),
        ("expected_source_url", f"https://github.com/{expected['source_repository']}/blob/{expected['source_commit']}/{expected['source_path']}"),
        ("expected_state", expected["state"]),
        ("promotion_gate", expected["promotion_gate"]),
        ("authority_transfer", str(expected["authority_transfer"]).lower()),
    ]
    for item in provenance:
        prov_ws.append(list(item))
    style_header(prov_ws)
    prov_ws.column_dimensions["A"].width = 34
    prov_ws.column_dimensions["B"].width = 120
    for row in prov_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(path)


def workbook_semantic_digest(path: Path) -> tuple[str, dict[str, list[list[Any]]]]:
    wb = load_workbook(path, data_only=False, read_only=True)
    semantic: dict[str, list[list[Any]]] = {}
    for ws in wb.worksheets:
        values = []
        for row in ws.iter_rows(values_only=True):
            values.append([v for v in row])
        semantic[ws.title] = values
    payload = json.dumps(
        semantic,
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return sha256_bytes(payload), semantic


def write_normalized_csv(rows: list[dict[str, str]], path: Path) -> None:
    headers = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def run() -> dict[str, Any]:
    baseline, rows, expected = load_inputs()
    guard_validation = validate_authority_guards(baseline, expected)
    validate_bindings(rows)
    calculated = calculate(baseline, rows)
    assert_expected(calculated, expected)

    OUT.mkdir(parents=True, exist_ok=True)
    workbook_path = OUT / "QPS_Partial_Relax_Workload_v1.xlsx"
    csv_path = OUT / "bindings_normalized.csv"
    receipt_path = OUT / "workload_receipt.json"

    write_workbook(baseline, rows, expected, calculated, workbook_path)
    write_normalized_csv(rows, csv_path)
    semantic_digest, semantic = workbook_semantic_digest(workbook_path)

    receipt = {
        "schema": "gbogeb.qps_rtm_partial_relax_workload/v1",
        "workload": "QPS_PARTIAL_RELAX_V06_EXPECTED_CALCULATION",
        "source_repository": expected["source_repository"],
        "source_commit": expected["source_commit"],
        "source_control_path": expected["source_path"],
        "source_control_blob_sha": expected["source_blob_sha"],
        "baseline_source_path": baseline["source_path"],
        "baseline_source_blob_sha": baseline["source_blob_sha"],
        "input_fixture_sha256": {
            "baseline_v05.json": sha256_file(BASELINE),
            "bindings.csv": sha256_file(BINDINGS),
            "expected_v06.json": sha256_file(EXPECTED),
        },
        "binding_rows": len(rows),
        "calculated": calculated,
        "expected_state": expected["state"],
        "guard_validation": guard_validation,
        "calculation_matches_expected_control": True,
        "excel_semantic_sha256": semantic_digest,
        "excel_sheet_names": list(semantic),
        "normalized_csv_sha256": sha256_file(csv_path),
        "status": "PASS_REPRODUCED_EXPECTED_V06_CALCULATION_NOT_PROMOTION",
        "authority_transfer": expected["authority_transfer"],
        "formal_credit_delta": expected["formal_credit_delta"],
        "engineering_credit_delta": expected["engineering_credit_delta"],
        "negotiation_credit_delta": expected["negotiation_credit_delta"],
        "compliance_credit_delta": expected["compliance_credit_delta"],
        "claim_guards": [
            "EXPECTED_V06_NE_PROMOTED_CURRENT_STATE",
            "WORKLOAD_REPRODUCIBILITY_NE_ENGINEERING_VALIDATION",
            "WORKLOAD_REPRODUCIBILITY_NE_CONTRACTUAL_OR_COMPLIANCE_CREDIT",
            "EXACT_V05_WORKBOOK_REGENERATION_GATE_REMAINS_SEPARATE",
        ],
    }
    receipt_path.write_text(json.dumps(receipt, indent=2) + "\n", encoding="utf-8")
    return receipt


if __name__ == "__main__":
    print(json.dumps(run(), sort_keys=True))
